from __future__ import annotations

from datetime import datetime, timezone
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from backend.app.errors import InvalidWorkbookError
from backend.app.knowledge_base.step_01_models import WorkbookSnapshot, WorkbookVersion
from backend.app.knowledge_base.step_02_loader import WorkbookLoader
from backend.app.knowledge_base.step_03_validator import WorkbookValidator
from backend.app.knowledge_base.step_04_service import KnowledgeBaseService
from backend.app.payloads.step_02_builder import PayloadBuilder
from backend.app.workflow.step_03_registry import WORKFLOW_HANDLER_METHODS
from backend.app.api import step_03_container
from backend.app.context.step_01_builder import GenerationContextBuilder
from backend.app.models.step_01_session import ContentSession

WORKBOOK = Path(
    os.getenv(
        "V2_TEST_WORKBOOK",
        "data/knowledge/test-client.xlsm",
    )
)


class KnowledgeBaseServiceTests(unittest.TestCase):
    def test_workbook_reload_does_not_initialize_the_full_service(self) -> None:
        with patch.object(step_03_container, "_services", {}):
            self.assertFalse(step_03_container.reload_v2_knowledge_if_initialized())

    def test_workbook_reload_refreshes_an_initialized_service(self) -> None:
        service = Mock()
        with patch.object(
            step_03_container,
            "_services",
            {step_03_container.get_active_client_id(): service},
        ):
            self.assertTrue(step_03_container.reload_v2_knowledge_if_initialized())
        service.knowledge.reload.assert_called_once_with()

    def _snapshot(self, sha256: str) -> WorkbookSnapshot:
        return WorkbookSnapshot(
            version=WorkbookVersion(
                filename="test.xlsm",
                sha256=sha256,
                loaded_at=datetime.now(timezone.utc),
                schema_version=None,
            ),
            post_types=(),
            taxonomies=(),
            shared_fields=(),
            acf_fields=(),
            blueprint=(),
            seo_rules=(),
            style_rules=(),
            story_patterns=(),
            image_analysis_rules=(),
            pillow_rules=(),
            image_metadata_fields=(),
            image_metadata_rules=(),
            internal_links=(),
            internal_link_rules=(),
            workflow_steps=(),
            application_states=(),
            agent_instructions=(),
            context_manifest=(),
            validation_values=(),
            post_examples=(),
            output_specifications=(),
        )

    def test_missing_workbook_hash_falls_back_to_current_snapshot(self) -> None:
        service = KnowledgeBaseService("unused.xlsm")
        current = self._snapshot("current-hash")
        service._snapshots[current.version.sha256] = current
        service._current_hash = current.version.sha256

        self.assertIs(service.by_hash("missing-hash"), current)

    def test_missing_workbook_hash_fallback_can_be_disabled(self) -> None:
        service = KnowledgeBaseService("unused.xlsm")
        current = self._snapshot("current-hash")
        service._snapshots[current.version.sha256] = current
        service._current_hash = current.version.sha256

        with patch.dict(os.environ, {"V2_ALLOW_CURRENT_WORKBOOK_FALLBACK": "0"}):
            with self.assertRaises(KeyError):
                service.by_hash("missing-hash")


class TaxonomyTermParsingTests(unittest.TestCase):
    def test_comma_separated_taxonomy_terms_become_a_list(self) -> None:
        self.assertEqual(
            PayloadBuilder._taxonomy_term_names(
                "Mojito, Negroni, Espresso Martini"
            ),
            ["Mojito", "Negroni", "Espresso Martini"],
        )

    def test_taxonomy_terms_are_trimmed_and_deduplicated(self) -> None:
        self.assertEqual(
            PayloadBuilder._taxonomy_term_names(
                ["Mojito; Negroni", "mojito", "", None]
            ),
            ["Mojito", "Negroni"],
        )


@unittest.skipUnless(WORKBOOK.is_file(), f"V2 test workbook not found: {WORKBOOK}")
class WorkbookTests(unittest.TestCase):
    def setUp(self) -> None:
        self.loader = WorkbookLoader()
        self.validator = WorkbookValidator()

    def test_v6_loads_and_all_exact_joins_validate(self) -> None:
        snapshot = self.validator.validate(self.loader.load(WORKBOOK))
        self.assertEqual(len(snapshot.version.sha256), 64)
        self.assertEqual(len(snapshot.workflow_steps), 20)
        self.assertEqual(len(snapshot.blueprint), 13)
        self.assertEqual(len(snapshot.image_metadata_rules), 2)
        self.assertEqual(
            {row.step_key for row in snapshot.workflow_steps},
            set(WORKFLOW_HANDLER_METHODS),
        )
        event_story = next(row for row in snapshot.acf_fields if row.field_key == "event_story")
        self.assertEqual((event_story.min_words, event_story.max_words), (80, 100))
        self.assertTrue(event_story.allow_internal_links)
        self.assertEqual(event_story.max_internal_links, 2)
        self.assertEqual(event_story.internal_link_priority, "high")

    def test_configured_fields_are_excluded_from_ai_schema(self) -> None:
        snapshot = self.validator.validate(self.loader.load(WORKBOOK))
        configured = [row for row in snapshot.shared_fields if row.source_mode == "configured"]
        self.assertTrue(configured)
        self.assertTrue(all(not row.include_in_ai_schema for row in configured))

    def test_post_shortcode_variables_load_as_input_fact_field_keys(self) -> None:
        snapshot = self.validator.validate(self.loader.load(WORKBOOK))
        bartender = snapshot.post_type("bartender")
        self.assertIsNotNone(bartender)
        self.assertEqual(bartender.post_shortcode_variables, ("staff_name",))
        staff_name = next(
            row
            for row in snapshot.acf_fields
            if row.post_type_key == "bartender" and row.field_key == "staff_name"
        )
        self.assertEqual(staff_name.field_role, "input_fact")

    def test_cocktail_knowledge_enrichment_and_instruction_load(self) -> None:
        snapshot = self.validator.validate(self.loader.load(WORKBOOK))
        cocktail = snapshot.post_type("cocktail")
        self.assertIsNotNone(cocktail)
        self.assertEqual(cocktail.knowledge_enrichment, "allowed")
        instruction = next(
            row
            for row in snapshot.agent_instructions
            if row.instruction_id == "client_003"
        )
        self.assertEqual(instruction.post_type_key, "cocktail")
        self.assertEqual(instruction.workflow_stage, "generation")
        self.assertEqual(instruction.priority, "critical")

    def test_unknown_knowledge_enrichment_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            copy = Path(temporary) / "bad-enrichment.xlsm"
            copy.write_bytes(WORKBOOK.read_bytes())
            workbook = load_workbook(copy, keep_vba=True)
            sheet = workbook["post_types"]
            post_type_column = next(
                cell.column for cell in sheet[1] if cell.value == "post_type_key"
            )
            enrichment_column = next(
                cell.column
                for cell in sheet[1]
                if cell.value == "knowledge_enrichment"
            )
            row = next(
                index
                for index in range(2, sheet.max_row + 1)
                if sheet.cell(index, post_type_column).value == "cocktail"
            )
            sheet.cell(row, enrichment_column).value = "unrestricted"
            workbook.save(copy)
            with self.assertRaises(InvalidWorkbookError) as raised:
                self.validator.validate(self.loader.load(copy))
            codes = {detail.error_code for detail in raised.exception.details}
            self.assertIn("unknown_knowledge_enrichment", codes)

    def test_bartender_taxonomy_relationship_loads_and_builds_payload(self) -> None:
        snapshot = self.validator.validate(self.loader.load(WORKBOOK))
        bartender = snapshot.post_type("bartender")
        self.assertIsNotNone(bartender)
        self.assertEqual(bartender.wp_taxonomies, ("barkeeper",))
        taxonomy = next(
            row for row in snapshot.taxonomies if row.wp_taxonomy == "barkeeper"
        )
        self.assertEqual(taxonomy.taxonomy_term_source, "acf:staff_name")
        self.assertTrue(taxonomy.assign_taxonomy_to_media)

        payload = PayloadBuilder().build(
            snapshot,
            post_type_key="bartender",
            shared_values={},
            acf_source_values={"staff_name": "Florent"},
        )
        self.assertEqual(payload.taxonomies, {"barkeeper": ["Florent"]})
        self.assertEqual(payload.media_taxonomies, ["barkeeper"])

    def test_post_type_can_select_several_taxonomies(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            copy = Path(temporary) / "several-taxonomies.xlsm"
            copy.write_bytes(WORKBOOK.read_bytes())
            workbook = load_workbook(copy, keep_vba=True)
            post_types = workbook["post_types"]
            taxonomies = workbook["taxonomies"]
            post_type_column = next(
                cell.column for cell in post_types[1] if cell.value == "post_type_key"
            )
            selections_column = next(
                cell.column for cell in post_types[1] if cell.value == "wp_taxonomies"
            )
            bartender_row = next(
                index
                for index in range(2, post_types.max_row + 1)
                if post_types.cell(index, post_type_column).value == "bartender"
            )
            post_types.cell(bartender_row, selections_column).value = "barkeeper, staff_role"
            taxonomies.append(("staff_role", "fixed:Barkeeper", False))
            workbook.save(copy)

            snapshot = self.validator.validate(self.loader.load(copy))
            payload = PayloadBuilder().build(
                snapshot,
                post_type_key="bartender",
                shared_values={},
                acf_source_values={"staff_name": "Florent"},
            )

        self.assertEqual(
            payload.taxonomies,
            {"barkeeper": ["Florent"], "staff_role": ["Barkeeper"]},
        )
        self.assertEqual(payload.media_taxonomies, ["barkeeper"])

    def test_same_taxonomy_can_use_a_source_specific_to_each_post_type(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            copy = Path(temporary) / "relationship-taxonomies.xlsm"
            copy.write_bytes(WORKBOOK.read_bytes())
            workbook = load_workbook(copy, keep_vba=True)
            post_types = workbook["post_types"]
            selections_column = next(
                cell.column for cell in post_types[1] if cell.value == "wp_taxonomies"
            )
            post_type_column = next(
                cell.column for cell in post_types[1] if cell.value == "post_type_key"
            )
            for row in range(2, post_types.max_row + 1):
                key = post_types.cell(row, post_type_column).value
                post_types.cell(row, selections_column).value = None
                if key in {"event", "bartender"}:
                    post_types.cell(row, selections_column).value = "barkeeper"

            del workbook["taxonomies"]
            taxonomies = workbook.create_sheet("taxonomies", 1)
            taxonomies.append(
                ("post_type_key", "wp_taxonomy", "taxonomy_term_source")
            )
            taxonomies.append(("bartender", "barkeeper", "acf:staff_name"))
            taxonomies.append(("event", "barkeeper", "acf:bartender"))
            workbook.save(copy)

            snapshot = self.validator.validate(self.loader.load(copy))
            event_payload = PayloadBuilder().build(
                snapshot,
                post_type_key="event",
                shared_values={},
                acf_source_values={"bartender": "Florent"},
            )
            bartender_payload = PayloadBuilder().build(
                snapshot,
                post_type_key="bartender",
                shared_values={},
                acf_source_values={"staff_name": "Alex"},
            )

        self.assertEqual(event_payload.taxonomies, {"barkeeper": ["Florent"]})
        self.assertEqual(bartender_payload.taxonomies, {"barkeeper": ["Alex"]})

    def test_reusable_prompt_guidance_resolves_for_selected_field(self) -> None:
        snapshot = self.validator.validate(self.loader.load(WORKBOOK))
        story_title = next(
            row
            for row in snapshot.acf_fields
            if row.post_type_key == "bartender" and row.field_key == "story_title"
        )
        self.assertEqual(
            story_title.prompt_guidance_keys,
            ("title_highlight_braces", "avoid_generic_intro", "concise_heading"),
        )
        session = ContentSession(
            session_id="session",
            user_id="user",
            post_type_key="bartender",
            state="ready_to_generate",
            workbook_hash=snapshot.version.sha256,
            language="de",
        )
        context = GenerationContextBuilder().build(
            snapshot=snapshot,
            task="content_generation",
            post_type_key="bartender",
            session=session,
            field_keys=["story_title"],
        )
        self.assertEqual(
            [row["guidance_key"] for row in context.fields["story_title"].prompt_guidance],
            ["title_highlight_braces", "avoid_generic_intro", "concise_heading"],
        )

    def test_unknown_reusable_prompt_guidance_fails_validation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            copy = Path(temporary) / "bad-prompt-guidance.xlsm"
            copy.write_bytes(WORKBOOK.read_bytes())
            workbook = load_workbook(copy, keep_vba=True)
            sheet = workbook["ACF_fields_schema"]
            field_column = next(cell.column for cell in sheet[1] if cell.value == "field_key")
            guidance_column = next(
                cell.column for cell in sheet[1] if cell.value == "prompt_guidance_keys"
            )
            row = next(
                index
                for index in range(2, sheet.max_row + 1)
                if sheet.cell(index, field_column).value == "story_title"
            )
            sheet.cell(row, guidance_column).value = "missing_guidance"
            workbook.save(copy)

            with self.assertRaises(InvalidWorkbookError) as raised:
                self.validator.validate(self.loader.load(copy))

            codes = {detail.error_code for detail in raised.exception.details}
            self.assertIn("unknown_prompt_guidance", codes)

    def test_invalid_post_shortcode_variable_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            copy = Path(temporary) / "bad-shortcode-variable.xlsm"
            copy.write_bytes(WORKBOOK.read_bytes())
            workbook = load_workbook(copy, keep_vba=True)
            sheet = workbook["post_types"]
            post_type_column = next(
                cell.column for cell in sheet[1] if cell.value == "post_type_key"
            )
            variables_column = next(
                cell.column
                for cell in sheet[1]
                if cell.value == "post shortcode variables"
            )
            row = next(
                index
                for index in range(2, sheet.max_row + 1)
                if sheet.cell(index, post_type_column).value == "bartender"
            )
            sheet.cell(row, variables_column).value = "not_an_input_fact"
            workbook.save(copy)
            with self.assertRaises(InvalidWorkbookError) as raised:
                self.validator.validate(self.loader.load(copy))
            codes = {detail.error_code for detail in raised.exception.details}
            self.assertIn("invalid_post_shortcode_variable", codes)

    def test_duplicate_active_url_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            copy = Path(temporary) / "duplicate.xlsm"
            copy.write_bytes(WORKBOOK.read_bytes())
            workbook = load_workbook(copy, keep_vba=True)
            sheet = workbook["internal_links_database"]
            target_url_column = next(
                cell.column for cell in sheet[1] if cell.value == "target_url"
            )
            active_column = next(cell.column for cell in sheet[1] if cell.value == "active")
            sheet.cell(3, target_url_column).value = sheet.cell(2, target_url_column).value
            sheet.cell(3, active_column).value = True
            workbook.save(copy)
            with self.assertRaises(InvalidWorkbookError) as raised:
                self.validator.validate(self.loader.load(copy))
            codes = {detail.error_code for detail in raised.exception.details}
            self.assertIn("duplicate_active_url", codes)

    def test_unknown_source_fact_key_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            copy = Path(temporary) / "bad-source.xlsm"
            copy.write_bytes(WORKBOOK.read_bytes())
            workbook = load_workbook(copy, keep_vba=True)
            sheet = workbook["ACF_fields_schema"]
            field_key_column = next(cell.column for cell in sheet[1] if cell.value == "field_key")
            source_column = next(cell.column for cell in sheet[1] if cell.value == "source_fact_keys")
            row = next(
                index
                for index in range(2, sheet.max_row + 1)
                if sheet.cell(index, field_key_column).value == "fact_event"
            )
            sheet.cell(row, source_column).value = "not_a_fact"
            workbook.save(copy)
            with self.assertRaises(InvalidWorkbookError) as raised:
                self.validator.validate(self.loader.load(copy))
            codes = {detail.error_code for detail in raised.exception.details}
            self.assertIn("unknown_source_fact_key", codes)
