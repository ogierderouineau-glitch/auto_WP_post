from __future__ import annotations

import hashlib
import logging
from datetime import datetime, timezone
from pathlib import Path
from threading import RLock
from typing import Any, TypeVar

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from backend.app.errors import ErrorDetail, InvalidWorkbookError
from backend.app.knowledge_base.app_owned import (
    APP_OWNED_AGENT_INSTRUCTIONS,
    APP_OWNED_APPLICATION_STATES,
    APP_OWNED_CONTEXT_MANIFEST,
    APP_OWNED_IMAGE_ANALYSIS_RULES,
    APP_OWNED_INTERNAL_LINK_RULES,
    APP_OWNED_OUTPUT_SPECIFICATIONS,
    APP_OWNED_PILLOW_RULES,
    APP_OWNED_WORKFLOW_STEPS,
    merge_agent_instructions,
)
from backend.app.knowledge_base.step_01_models import (
    ACFFieldSchema,
    AgentInstruction,
    BlueprintRow,
    HTMLPattern,
    ImageMetadataField,
    ImageMetadataRule,
    InternalLinkRecord,
    PostExample,
    PostTypeConfig,
    PromptGuidance,
    SEORule,
    SharedFieldSchema,
    StoryPattern,
    StyleRule,
    TaxonomyConfig,
    ValidationListValue,
    WorkbookSnapshot,
    WorkbookVersion,
)

ModelT = TypeVar("ModelT", bound=BaseModel)
LOGGER = logging.getLogger(__name__)

REQUIRED_SHEETS = {
    "post_types",
    "shared_fields_schema",
    "seo_rules",
    "ACF_fields_schema",
    "post_blueprint",
    "story_patterns",
    "style_rules",
    "image_metadata_schema",
    "image_metadata_rules",
    "internal_links_database",
    "validation_lists",
    "post_examples",
    "agent_instructions",
}

DEPRECATED_APP_OWNED_SHEETS = {
    "agent_workflow",
    "application_state",
    "context_building",
    "image_analysis_rules",
    "image_rules_pillow",
    "internal_link_rules",
    "output_specification",
    "agent_instructions_app",
}

LIST_COLUMNS = {
    "source_fact_keys",
    "context_tags",
    "trigger_tags",
    "required_facts",
    "required_content_signals",
    "excluded_tags",
    "semantic_prompt_hints",
    "anchor_variants",
    "trigger_value",
    "target_field_keys",
    "allowed_next_states",
    "required_data",
    "post shortcode variables",
    "post_shortcode_variables",
    "wp_taxonomies",
    "html_pattern_keys",
    "prompt_guidance_keys",
    "allowed_field_keys",
}

BOOLEAN_COLUMNS = {
    "enabled",
    "approved",
    "required",
    "active",
    "template_ready",
    "generation_enabled",
    "user_selectable",
    "required_for_output",
    "required_for_analysis",
    "include_in_ai_schema",
    "include_in_payload",
    "include_in_image_metadata_context",
    "allow_internal_links",
    "terminal",
}


def _split_list(value: Any) -> tuple[str, ...]:
    if value is None:
        return ()
    return tuple(part.strip() for part in str(value).split(";") if part.strip())


def _parse_boolean(value: Any, *, sheet: str, row: int, column: str) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value or "").strip().upper()
    if normalized == "TRUE":
        return True
    if normalized == "FALSE":
        return False
    raise InvalidWorkbookError(
        "Workbook contains an invalid boolean.",
        details=[
            ErrorDetail(
                sheet=sheet,
                row=row,
                column=column,
                error_code="invalid_boolean",
                message=f"Expected TRUE/FALSE or an Excel boolean, received {value!r}.",
            )
        ],
    )


class WorkbookLoader:
    """Load an immutable, typed workbook snapshot and cache it by SHA-256."""

    def __init__(self) -> None:
        self._cache: dict[str, WorkbookSnapshot] = {}
        self._lock = RLock()

    def load(self, path: str | Path) -> WorkbookSnapshot:
        workbook_path = Path(path).expanduser().resolve()
        if not workbook_path.is_file():
            raise InvalidWorkbookError(
                f"V2 workbook not found: {workbook_path}",
                details=[
                    ErrorDetail(
                        sheet=None,
                        row=None,
                        column=None,
                        error_code="workbook_not_found",
                        message=str(workbook_path),
                    )
                ],
            )
        digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
        with self._lock:
            cached = self._cache.get(digest)
            if cached is not None:
                return cached

        workbook = load_workbook(workbook_path, read_only=False, data_only=True, keep_vba=True)
        missing = sorted(REQUIRED_SHEETS.difference(workbook.sheetnames))
        if missing:
            raise InvalidWorkbookError(
                "Workbook is missing required sheets.",
                details=[
                    ErrorDetail(
                        sheet=name,
                        row=None,
                        column=None,
                        error_code="missing_required_sheet",
                        message=f"Required sheet {name!r} is missing.",
                    )
                    for name in missing
                ],
            )
        self._warn_deprecated_sheets(workbook)
        post_types, taxonomies = self._post_types_and_taxonomies(workbook)

        snapshot = WorkbookSnapshot(
            version=WorkbookVersion(
                filename=workbook_path.name,
                sha256=digest,
                loaded_at=datetime.now(timezone.utc),
                schema_version=self._schema_version(workbook),
            ),
            post_types=post_types,
            taxonomies=taxonomies,
            shared_fields=self._models(workbook["shared_fields_schema"], SharedFieldSchema),
            acf_fields=self._models(workbook["ACF_fields_schema"], ACFFieldSchema),
            prompt_guidance=(
                self._models(workbook["prompt_guidance"], PromptGuidance)
                if "prompt_guidance" in workbook.sheetnames
                else ()
            ),
            html_patterns=(
                self._models(workbook["HTML_patterns"], HTMLPattern)
                if "HTML_patterns" in workbook.sheetnames
                else ()
            ),
            blueprint=self._models(workbook["post_blueprint"], BlueprintRow),
            seo_rules=self._models(workbook["seo_rules"], SEORule),
            style_rules=self._models(workbook["style_rules"], StyleRule),
            story_patterns=self._models(workbook["story_patterns"], StoryPattern),
            image_analysis_rules=APP_OWNED_IMAGE_ANALYSIS_RULES,
            pillow_rules=APP_OWNED_PILLOW_RULES,
            image_metadata_fields=self._models(workbook["image_metadata_schema"], ImageMetadataField),
            image_metadata_rules=self._models(workbook["image_metadata_rules"], ImageMetadataRule),
            internal_links=self._models(workbook["internal_links_database"], InternalLinkRecord),
            internal_link_rules=APP_OWNED_INTERNAL_LINK_RULES,
            workflow_steps=APP_OWNED_WORKFLOW_STEPS,
            application_states=APP_OWNED_APPLICATION_STATES,
            agent_instructions=merge_agent_instructions(
                APP_OWNED_AGENT_INSTRUCTIONS,
                self._client_agent_instructions(workbook),
            ),
            context_manifest=APP_OWNED_CONTEXT_MANIFEST,
            validation_values=self._models(workbook["validation_lists"], ValidationListValue),
            post_examples=self._models(workbook["post_examples"], PostExample),
            output_specifications=APP_OWNED_OUTPUT_SPECIFICATIONS,
        )
        with self._lock:
            self._cache[digest] = snapshot
        return snapshot

    def _post_types_and_taxonomies(
        self,
        workbook: Any,
    ) -> tuple[tuple[PostTypeConfig, ...], tuple[TaxonomyConfig, ...]]:
        post_types = self._models(workbook["post_types"], PostTypeConfig)
        if "taxonomies" in workbook.sheetnames:
            return post_types, self._models(workbook["taxonomies"], TaxonomyConfig)

        worksheet = workbook["post_types"]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        legacy_columns = {
            "wp_taxonomy",
            "taxonomy_term_source",
            "assign_taxonomy_to_media",
        }
        if not legacy_columns.issubset(headers):
            raise InvalidWorkbookError(
                "Workbook is missing the taxonomies sheet.",
                details=[
                    ErrorDetail(
                        sheet="taxonomies",
                        row=None,
                        column=None,
                        error_code="missing_required_sheet",
                        message=(
                            "Add a taxonomies sheet, or retain the legacy taxonomy "
                            "columns on post_types during migration."
                        ),
                    )
                ],
            )

        migrated_post_types: list[PostTypeConfig] = []
        taxonomies: list[TaxonomyConfig] = []
        for post_type in post_types:
            row = post_type.sheet_row
            slug = str(
                worksheet.cell(row, headers["wp_taxonomy"]).value or ""
            ).strip().lower()
            source = str(
                worksheet.cell(row, headers["taxonomy_term_source"]).value or ""
            ).strip()
            assign_media = worksheet.cell(
                row,
                headers["assign_taxonomy_to_media"],
            ).value
            migrated_post_types.append(
                post_type.model_copy(
                    update={"wp_taxonomies": (slug,) if slug else ()}
                )
            )
            if slug or source or assign_media:
                taxonomies.append(TaxonomyConfig.model_validate({
                    "sheet_row": row,
                    "post_type_key": post_type.post_type_key,
                    "wp_taxonomy": slug,
                    "taxonomy_term_source": source,
                    "assign_taxonomy_to_media": bool(assign_media),
                }))
        return tuple(migrated_post_types), tuple(taxonomies)

    @staticmethod
    def _schema_version(workbook: Any) -> str | None:
        if "README" not in workbook.sheetnames:
            return None
        for row in workbook["README"].iter_rows(values_only=True):
            if row and str(row[0] or "").strip().lower() in {"schema_version", "schema version"}:
                return str(row[1] or "").strip() or None
        return None

    def _client_agent_instructions(self, workbook: Any) -> tuple[AgentInstruction, ...]:
        sheet_names = ["agent_instructions"]
        secondary_names = sorted(
            name
            for name in workbook.sheetnames
            if name.startswith("agent_instructions")
            and name not in {"agent_instructions", "agent_instructions_app"}
        )
        sheet_names.extend(secondary_names)
        rules: list[AgentInstruction] = []
        for sheet_name in sheet_names:
            if sheet_name in workbook.sheetnames:
                rules.extend(self._models(workbook[sheet_name], AgentInstruction))
        return tuple(rules)

    @staticmethod
    def _warn_deprecated_sheets(workbook: Any) -> None:
        present = sorted(DEPRECATED_APP_OWNED_SHEETS.intersection(workbook.sheetnames))
        if not present:
            return
        LOGGER.warning(
            "Workbook contains deprecated app-owned sheets that are ignored: %s",
            ", ".join(present),
        )

    def _models(self, worksheet: Any, model: type[ModelT]) -> tuple[ModelT, ...]:
        headers = [str(cell.value).strip() if cell.value is not None else None for cell in worksheet[1]]
        required_headers = {
            name
            for name, field in model.model_fields.items()
            if name != "sheet_row" and field.is_required()
        }
        missing_headers = sorted(required_headers.difference(header for header in headers if header))
        if missing_headers:
            raise InvalidWorkbookError(
                f"Sheet {worksheet.title!r} is missing required columns.",
                details=[
                    ErrorDetail(
                        sheet=worksheet.title,
                        row=1,
                        column=column,
                        error_code="missing_required_column",
                        message=f"Required column {column!r} is missing.",
                    )
                    for column in missing_headers
                ],
            )

        parsed: list[ModelT] = []
        for row_number in range(2, worksheet.max_row + 1):
            raw = {
                header: worksheet.cell(row_number, column_number).value
                for column_number, header in enumerate(headers, 1)
                if header
            }
            if not any(value is not None for value in raw.values()):
                continue
            normalized = self._normalize(raw, worksheet.title, row_number)
            normalized["sheet_row"] = row_number
            try:
                parsed.append(model.model_validate(normalized))
            except ValidationError as exc:
                raise InvalidWorkbookError(
                    f"Sheet {worksheet.title!r} contains invalid typed data.",
                    details=[
                        ErrorDetail(
                            sheet=worksheet.title,
                            row=row_number,
                            column=".".join(str(part) for part in issue["loc"]),
                            error_code="invalid_typed_value",
                            message=issue["msg"],
                            context={"input": issue.get("input")},
                        )
                        for issue in exc.errors()
                    ],
                ) from exc
        return tuple(parsed)

    @staticmethod
    def _normalize(raw: dict[str, Any], sheet: str, row: int) -> dict[str, Any]:
        normalized = dict(raw)
        for column in LIST_COLUMNS.intersection(normalized):
            normalized[column] = (
                tuple(
                    part.strip()
                    for part in str(normalized[column] or "").replace(",", ";").split(";")
                    if part.strip()
                )
                if column in {"prompt_guidance_keys", "wp_taxonomies"}
                else _split_list(normalized[column])
            )
        for column in BOOLEAN_COLUMNS.intersection(normalized):
            if normalized[column] is not None:
                normalized[column] = _parse_boolean(
                    normalized[column],
                    sheet=sheet,
                    row=row,
                    column=column,
                )
        if "example_id" in normalized and normalized["example_id"] is not None:
            normalized["example_id"] = str(normalized["example_id"]).removesuffix(".0")
        return normalized
